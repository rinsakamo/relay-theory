#!/usr/bin/env python3
"""P20 main source-critical original publisher figure/table extensions. Never MAIN science."""
import io,hashlib,json
from pathlib import Path
from urllib.request import Request,urlopen
from PIL import Image
from docx import Document
out=Path("g1-w4-p20-supplements");out.mkdir(exist_ok=True)
rows=[];bad=[]
for i in range(2,10):
    kind="TIFF" if i<=5 else "DOCX"
    name=f"P20_S{i-1}_Fig" if i<=5 else f"P20_S{i-5}_Table"
    doi=f"10.1371/journal.pcbi.1006676.s{i:03d}"
    url=f"https://journals.plos.org/ploscompbiol/article/file?id={doi}&type=supplementary"
    rec={"name":name,"doi":doi,"url":url,"type_expected":kind}
    try:
        res=urlopen(Request(url,headers={"User-Agent":"RelayTheory-W4-source-fidelity/1"}),timeout=65)
        raw=res.read();rec.update(raw_sha256=hashlib.sha256(raw).hexdigest(),bytes=len(raw),publisher_status="FETCHED")
        if kind=="TIFF":
            im=Image.open(io.BytesIO(raw))
            if im.format!="TIFF":raise ValueError(f"expected TIFF actual {im.format}")
            rec["dimensions"]=list(im.size)
            (out/(name+".tif")).write_bytes(raw)
        else:
            if not raw.startswith(b"PK"):raise ValueError("DOCX container signature mismatch")
            doc=Document(io.BytesIO(raw))
            rec["table_count"]=len(doc.tables)
            rec["text_excerpts"]=[p.text for p in doc.paragraphs if p.text.strip()][:5]
            rec["table_preview"]=[[[c.text for c in row.cells] for row in tab.rows[:3]] for tab in doc.tables[:2]]
            (out/(name+".docx")).write_bytes(raw)
    except Exception as e:
        rec.update(publisher_status="FAIL",error=repr(e));bad.append(name+":"+repr(e))
    rows.append(rec)
    print(json.dumps(rec,ensure_ascii=False),flush=True)
(out/"source_receipt.json").write_text(json.dumps({"original_source_status":"PRE_A_ONLY","rows":rows,"errors":bad},ensure_ascii=False,sort_keys=True,indent=2)+"\n")
if bad:raise SystemExit("Official P20 defining source blocked: "+repr(bad))
