#!/usr/bin/env python3
"""Separate real P20 after-E original publisher eleven exact SHA source + immutable stage history verifier.
Exact user data/agent numerical replay and independent human semantic review NOT performed.
"""
import io,json,hashlib,subprocess,zipfile
from pathlib import Path
from urllib.request import Request,urlopen
from pypdf import PdfReader
from PIL import Image
from docx import Document
import xml.etree.ElementTree as ET
S="research/paper2/p399/g1/parallel/W4/p20/"
STAGES=[
("PRE_A","6eed47b516114455d9af8fe69553de8c02001b3f","P20_PRE_A_COMPLETE_ORIGINAL_MAIN_S1_S10_FREEZE.json"),
("A","aa5e64fb080870ae81f7deb8428850ae93e4f8a1","P20_A_COMPLETE_SOURCE_36_CLAIMS_46_LINKS.json"),
("B","e24c848c06f4a96700d1705f796d423b4ec61a36","P20_B_COMPLETE_36_47_ORIGINAL_SOURCE_CHALLENGE.json"),
("C","767d73b8ea868efe150c621e03c9638b3d6f0be6","P20_C_FULL_C1C2_ADJUDICATION_36_47.json"),
("D","558fa03e4ceb34df6a1982aa92cab43103574b42","P20_D_FULL_SOURCE_36_47_UNCHANGED_GRAMMAR_V0.json"),
("E","6b6b4cb685a1747ece22af46972286c2226eca13","P20_E_COMPLETE_SOURCE_FIDELITY_10_TARGETS.json")]
SRCS={
"main":("pdf",3073632,27,"5b6c3347eda5fc1509db0867e2c75fe94e90f1fe85385abf61f8738b2f16f097"),
"s001":("zip",2326772,None,"06caca17f908fa58dff9ae2fe19bbf07e16092e333b38342287d2429f9b89103"),
"s002":("tiff",1732668,None,"e1e1e268d7c00cb5afcc3b01a51474519a306dcd009b0d94dcc9aa881d7c9589"),
"s003":("tiff",835636,None,"90203284187c28581356885db7baf7ccbf644ab0c67242e37215147dec04c6bf"),
"s004":("tiff",671194,None,"066f93043104b855043dd35f6be9c457ff0c237e0cb4669f4675fabc6c8b0c2b"),
"s005":("tiff",799556,None,"1c734f7c85abd7293de9bb8140f9c6339408fbc5c30870c9d07a719adb7233f6"),
"s006":("docx",12614,None,"63f8be97490aec5ebd7c215f0efd6c1a6eb7fc68e0ebd808222f16cf4a75f482"),
"s007":("docx",12731,None,"b0621905ad760016d827b7424276a50634f2bd916ff196357dc2106c04cca2a0"),
"s008":("docx",12744,None,"0d217f48cc8deaf8e6446803e3ce9903c0c55d77eded998af702ce9e90e76d67"),
"s009":("docx",12224,None,"0651f92e8ee222bbdd10f5e57f7d0f0831856f928e0a0bffc50f37291afb5c0f"),
"s010":("mp4",3002196,None,"2fb702c0f629f8f9e2288adbdf96a17a1db6b68b97cea21006f2298673b5c7f0")
}
out=Path("g1-w4-p20-postE");out.mkdir(exist_ok=True)
records=[];bad=[];sha=lambda bs:hashlib.sha256(bs).hexdigest()
def audit(k,pass_,detail):
 item={"check":k,"pass":bool(pass_),"details":detail};records.append(item)
 if not pass_:bad.append(item)
 print(json.dumps(item,sort_keys=True,ensure_ascii=False),flush=True)
def git(*args):return subprocess.check_output(["git",*args])
try:
 last="38c5eb24c61b9e57371084a65fcbcb9fb9c62a4b"
 for n,commit,fname in STAGES:
  name=S+fname
  old=git("show",commit+":"+name);new=Path(name).read_bytes()
  p=subprocess.run(["git","merge-base","--is-ancestor",last,commit],capture_output=True)
  audit("frozen_stage_original_byte_and_ancestry_"+n,sha(old)==sha(new) and p.returncode==0,{"commit":commit,"old_raw_sha":sha(old),"head_sha":sha(new),"parent_stage_commit":last})
  last=commit
 j=json.loads(Path(S+"P20_A_COMPLETE_SOURCE_36_CLAIMS_46_LINKS.json").read_text())
 audit("original_A_actual_36_47_despite_46_filename",len(j["claims"])==36 and len(j["dependencies"])==47 and j["dependency_count"]==47,{"claims":len(j["claims"]),"source_dependencies":len(j["dependencies"])})
except Exception as e:audit("historic_provenance_exception",False,repr(e))
try:
 xml=urlopen("https://www.ebi.ac.uk/europepmc/webservices/rest/PMC6420027/fullTextXML",timeout=70).read()
 root=ET.fromstring(xml)
 ids=[x.text for x in root.findall(".//article-id") if x.attrib.get("pub-id-type")=="doi"]
 audit("independent_same_issuer_deposited_jats_original_doi","10.1371/journal.pcbi.1006676" in ids,{"dois":ids,"jats_sha":sha(xml)})
 zipbs=urlopen("https://www.ebi.ac.uk/europepmc/webservices/rest/PMC6420027/supplementaryFiles",timeout=95).read()
 z=zipfile.ZipFile(io.BytesIO(zipbs)); archive_sha={}
 for key,(ext,size,pages,expect) in SRCS.items():
  if key=="main":continue
  file="pcbi.1006676."+key+"."+ext
  bs=z.read(file)
  archive_sha[key]=sha(bs)
  audit("NLM_original_archived_"+key,sha(bs)==expect and len(bs)==size,{"member":file,"sha":sha(bs),"bytes":len(bs)})
except Exception as e:
 audit("reacquire_original_DOI_archived_bundle_exception",False,repr(e));archive_sha={}
for key,(ext,n,pages,expect) in SRCS.items():
 name="10.1371/journal.pcbi.1006676"+("" if key=="main" else "."+key)
 url="https://journals.plos.org/ploscompbiol/article/file?id="+name+"&type="+("printable" if key=="main" else "supplementary")
 try:
  bs=urlopen(Request(url,headers={"User-Agent":"RelayTheory-W4-independent-POST-E-original-verifier/1"}),timeout=65).read()
  rec={"publisher":name,"raw_sha":sha(bs),"bytes":len(bs)}
  ok=sha(bs)==expect and len(bs)==n
  if key=="main":
   pdf=PdfReader(io.BytesIO(bs));rec["pages"]=len(pdf.pages);ok=ok and len(pdf.pages)==27
   anchors={1:"Goal-related feedback",8:"Participants",19:"forward kinematics"}
   for page,t in anchors.items():
    found=t.lower() in " ".join((pdf.pages[page-1].extract_text() or "").split()).lower()
    audit("P20_main_original_text_page_"+str(page),found,{"page":page,"term":t})
  elif ext=="tiff":
   img=Image.open(io.BytesIO(bs));rec["format"]=img.format;rec["dimensions"]=list(img.size);ok=ok and img.format=="TIFF"
  elif ext=="docx":
   d=Document(io.BytesIO(bs))
   rec["nonempty_source_paragraphs"]=sum(bool(x.text.strip()) for x in d.paragraphs)
   if key=="s006":
    text=" ".join(x.text for x in d.paragraphs)
    audit("original_P20_S1_Table_null_H1_H2_text_anchor","0.944" in text or ".944" in text,{"anchor":"source p=.944"})
  elif ext=="zip":
   zz=zipfile.ZipFile(io.BytesIO(bs));rec["archive_members"]=len(zz.namelist());ok=ok and len(zz.namelist())>=220
  elif ext=="mp4":
   rec["file_signature"]="ftyp" in str(bs[:16]);ok=ok and rec["file_signature"]
  if key in archive_sha:ok=ok and sha(bs)==archive_sha[key]
  audit("POST_E_current_live_issuer_same_original_"+key,ok,rec)
 except Exception as e:audit("POST_E_CURRENT_ISSUER_unavailable_"+key,False,repr(e))
result={"mode":"INDEPENDENT_REAL_AFTER_E_FROZEN_SOURCE_AND_GIT_NOT_NUMERIC_REPLAY","checks":records,"failures":bad,"PASS":not bad}
(out/"P20_AFTER_E_RAW_SHA_PAGE_ANCHORS_GIT.json").write_text(json.dumps(result,ensure_ascii=False,sort_keys=True,indent=2)+"\n")
print("W4_P20_POST_E_FAILURES",len(bad),flush=True)
if bad:raise SystemExit(1)
